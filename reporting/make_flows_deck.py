"""Build the algorithm-flow deck from the figures in reporting/charts/flows.

One question per slide, the figure that answers it, and the number the figure is
there to support. Nothing on a slide that is not derived from
out/analysis/algorithm_flows.json or the ledger behind it.

    python reporting/make_flows_deck.py [--out out/analysis/dnssec_algorithm_flows.pptx]
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

W, H = Inches(13.333), Inches(7.5)
INK, INK_2, MUTED = RGBColor(0x0B, 0x0B, 0x0B), RGBColor(0x52, 0x51, 0x4E), RGBColor(0x89, 0x87, 0x81)
ACCENT = RGBColor(0x2A, 0x78, 0xD6)
SURFACE = RGBColor(0xFC, 0xFC, 0xFB)
FONT = "Segoe UI"


def text(slide, x, y, w, h, runs, size=18, color=INK, bold=False, space=6):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space)
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.name = FONT
        r.font.color.rgb = color
        r.font.bold = bold
    return box


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = SURFACE
    return s


def chart_slide(prs, kicker, title, image, takeaway, caption):
    s = blank(prs)
    text(s, Inches(0.62), Inches(0.38), Inches(12), Inches(0.3), [kicker.upper()],
         size=12, color=MUTED, bold=True)
    text(s, Inches(0.62), Inches(0.68), Inches(12.1), Inches(0.7), [title],
         size=28, bold=True)
    # Size by height: these figures are ~2.4:1, so fitting to width overruns the
    # takeaway line at 6.28".
    band_top, band_h = Inches(1.58), Inches(4.52)
    pic = s.shapes.add_picture(str(image), Inches(0.9), band_top, height=band_h)
    if pic.width > Inches(11.9):
        pic.width, pic.height = Inches(11.9), int(pic.height * Inches(11.9) / pic.width)
    pic.left = int((W - pic.width) / 2)
    pic.top = int(band_top + (band_h - pic.height) / 2)   # centre in the band
    text(s, Inches(0.62), Inches(6.28), Inches(12.1), Inches(0.45), [takeaway],
         size=16, color=ACCENT, bold=True)
    text(s, Inches(0.62), Inches(6.78), Inches(12.1), Inches(0.5), [caption],
         size=11, color=MUTED)
    return s


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--figs", type=Path, default=Path("reporting/charts/flows/bare"),
                    help="title-less figures; the slide supplies the wording")
    ap.add_argument("--data", type=Path, default=Path("out/analysis/algorithm_flows.json"))
    ap.add_argument("--out", type=Path,
                    default=Path("out/analysis/dnssec_algorithm_flows.pptx"))
    args = ap.parse_args()

    d = json.loads(args.data.read_text(encoding="utf-8"))
    share = d["first_signing_share_pct"]
    base = d["installed_base_last"]
    ecdsa_2026 = share["2026"]["ECDSA"]
    first_half = next(y for y in sorted(share) if (share[y]["ECDSA"] or 0) > 50)

    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    # ---- title ------------------------------------------------------------- #
    s = blank(prs)
    text(s, Inches(0.9), Inches(2.3), Inches(11.5), Inches(0.4),
         ["REVERSE-DNS DELEGATIONS, 2009-2026"], size=13, color=MUTED, bold=True)
    text(s, Inches(0.9), Inches(2.75), Inches(11.6), Inches(1.6),
         ["What DNSSEC zones are switching to,", "and what they are leaving"],
         size=40, bold=True, space=2)
    line = s.shapes.add_shape(1, Inches(0.92), Inches(4.5), Inches(2.0), Emu(28575))
    line.fill.solid(); line.fill.fore_color.rgb = ACCENT; line.line.fill.background()
    text(s, Inches(0.9), Inches(4.9), Inches(11.5), Inches(1.4), [
        "25,930 algorithm changes across 13,654 delegations, tracked zone by zone.",
        "Two complete generational handovers: RSA/SHA-1 to RSA/SHA-2 to ECDSA.",
        "Reverse corpus only — the forward per-day records are not in this repository.",
    ], size=15, color=INK_2)

    # ---- the five figures -------------------------------------------------- #
    chart_slide(
        prs, "the choice", "What a zone turning DNSSEC on picks",
        args.figs / "f1_first_signing_choice.png",
        f"ECDSA is now {ecdsa_2026:.0f}% of new signings; it was 0% before 2015.",
        "Zones gaining a DS record that year, by algorithm family. 2026 is nine months.")

    chart_slide(
        prs, "the handover", "Two complete handovers in seventeen years",
        args.figs / "f2_first_signing_share.png",
        f"ECDSA passes half of new signings in {first_half} and reaches "
        f"{ecdsa_2026:.0f}% by 2026.",
        "Same data as a share of each year's total. EdDSA, standardised in 2017, "
        "never becomes visible at this scale.")

    chart_slide(
        prs, "the migration", "Where zones that already had DNSSEC moved",
        args.figs / "f3_rollover_paths.png",
        "One path dominates: 403 zones went RSASHA256 to ECDSA P-256.",
        "\"added X\" is a dual-signing step, publishing the new algorithm alongside the "
        "old one — how a correct rollover begins. A quarter of rollovers pass through it.")

    chart_slide(
        prs, "arrivals and departures", "Each family's flow, year by year",
        args.figs / "f4_arrivals_departures.png",
        "ECDSA has never had a year of net departures. RSA/SHA-1 has had eight.",
        "A zone arrives in the year it first publishes an algorithm from that family and "
        "leaves in the year it stops.")

    chart_slide(
        prs, "the stock", "The installed base those flows add up to",
        args.figs / "f5_installed_base.png",
        f"ECDSA {base['ECDSA']:,} · RSA/SHA-2 {base['RSA/SHA-2']:,} · "
        f"RSA/SHA-1 {base['RSA/SHA-1']:,} · EdDSA {base['EdDSA']:,}",
        "Running total of arrivals minus departures. A net position over observed "
        "changes, not a census of the reverse DNS.")

    # ---- what it means ----------------------------------------------------- #
    s = blank(prs)
    text(s, Inches(0.62), Inches(0.38), Inches(12), Inches(0.3), ["WHAT THIS SHOWS"],
         size=12, color=MUTED, bold=True)
    text(s, Inches(0.62), Inches(0.68), Inches(12.1), Inches(0.7),
         ["Three things worth taking away"], size=28, bold=True)
    items = [
        ("Algorithm choice turns over completely, about once a decade.",
         "RSA/SHA-1 was the usual choice until 2013, RSA/SHA-2 until 2020, ECDSA since. "
         "From first appearance to majority of new signings took RSA/SHA-2 four years "
         "(2010 to 2014) and ECDSA six (2015 to 2021)."),
        ("Rollovers are rare, and one path carries most of them.",
         "1,551 rollovers against 18,205 first-time signings. Of those, 403 are RSASHA256 "
         "to ECDSA P-256 — more than the next two paths combined."),
        ("EdDSA has not started.",
         f"Standardised in 2017. Thirty-eight zones in the net position, against "
         f"{base['ECDSA']:,} for ECDSA. Nine years is long enough that this is a result, "
         "not a wait."),
    ]
    y = Inches(1.7)
    for i, (head, body) in enumerate(items, 1):
        text(s, Inches(0.62), y, Inches(0.5), Inches(0.5), [str(i)], size=24,
             color=ACCENT, bold=True)
        text(s, Inches(1.25), y, Inches(11.4), Inches(0.4), [head], size=18, bold=True)
        text(s, Inches(1.25), y + Inches(0.42), Inches(11.4), Inches(0.8), [body],
             size=13, color=INK_2)
        y += Inches(1.55)
    text(s, Inches(0.62), Inches(6.9), Inches(12), Inches(0.3),
         ["Reproducible: python reporting/algorithm_flows.py && "
          "python reporting/make_flows_deck.py"], size=10, color=MUTED)

    args.out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(args.out)
    print(f"wrote {args.out}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
